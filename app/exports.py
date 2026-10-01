from io import BytesIO
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn


def sections(meeting, result):
    yield "Резюме", [result["summary"]]
    yield "Участники", result["participants"] or ["Не указаны"]
    yield "Темы", result["topics"] or ["Не указаны"]
    yield "Решения", [
        f"{d['text']}\nИсточник: {d['quote']}" for d in result["decisions"]
    ] or ["Решения не зафиксированы"]
    yield "Поручения", [
        f"{t['text']}\nОтветственный: {t['owner'] or 'Не указано'}\nСрок: {t['due_date'] or 'Не указано'}\nСтатус: {t['status']}\nИсточник: {t['quote']}"
        for t in result["tasks"]
    ] or ["Поручения не зафиксированы"]
    yield "Открытые вопросы", result["open_questions"] or ["Не зафиксированы"]


def markdown(meeting, result):
    parts = [f"# {meeting['title']}", f"Дата встречи: {meeting['meeting_date']}"]
    for heading, lines in sections(meeting, result):
        parts.append("## " + heading)
        parts.extend(lines)
    return "\n\n".join(parts) + "\n"


def docx_bytes(meeting, result):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "DejaVu Sans"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.1
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(2)
        section.left_margin = section.right_margin = Cm(2)
    for style in doc.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    doc.styles["Title"].font.size = Pt(22)
    doc.styles["Heading 1"].font.size = Pt(14)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(10)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(4)
    for name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    doc.add_heading(meeting["title"], 0)
    doc.add_paragraph("Дата встречи: " + meeting["meeting_date"])
    for heading, lines in sections(meeting, result):
        doc.add_heading(heading, 1)
        for line in lines:
            doc.add_paragraph(line)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()
